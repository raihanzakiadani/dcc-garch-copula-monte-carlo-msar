import json, re
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
CFG = json.loads((ROOT / "config/pipeline.json").read_text())
GS = (ROOT / "apps_script/yahoo_feed.gs").read_text()


def test_apps_script_column_labels_equal_bloomberg_column_names():
    pairs = re.findall(r"\{col: '([^']+)',\s*symbol: '([^']+)'\}", GS)
    labels = {c for c, _ in pairs}
    in_cfg = {k for k, v in CFG["instruments"].items() if v.get("csv_url_env") == "SRC_YAHOO_SHEET_CSV"}
    assert labels == in_cfg == {"USDIDR BGN Curncy", "CO1 Comdty"}
    assert dict(pairs) == {"USDIDR BGN Curncy": "IDR=X", "CO1 Comdty": "BZ=F"}


def test_every_config_column_exists_in_the_bloomberg_dataset_header():
    from openpyxl import load_workbook
    hdr = [c for c in next(load_workbook(ROOT / "seed/raw_data.xlsx", read_only=True).worksheets[0].iter_rows(values_only=True))]
    assert set(CFG["instruments"]) <= set(hdr)
    for v in CFG["instruments"].values():
        assert v.get("column", "x") not in ("",)          # sheet column label must be non-empty when set


def test_doc_table_lists_every_instrument():
    doc = (ROOT / "docs/TICKER_MAP.md").read_text()
    for k in CFG["instruments"]:
        assert f"`{k}`" in doc, k
