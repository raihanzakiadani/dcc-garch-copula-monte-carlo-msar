"""One-off: remove duplicate dates from raw_data.xlsx (policy: keep the LAST occurrence = later paste supersedes).
Only deletes rows; every other cell (including literal 'NA' text) is left exactly as it was.
  python scripts/clean_raw_data.py seed/raw_data.xlsx --report docs/clean_report.json
"""
import argparse, datetime as dt, json
from openpyxl import load_workbook

ap = argparse.ArgumentParser(); ap.add_argument("xlsx"); ap.add_argument("--report", default=None)
a = ap.parse_args()
wb = load_workbook(a.xlsx); ws = wb.worksheets[0]
seen, drop = {}, []
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, 1).value
    d = v.date() if isinstance(v, dt.datetime) else v
    if d is None: continue
    if d in seen: drop.append(seen[d])
    seen[d] = r
drop = sorted(drop, reverse=True)
rep = {"rows_before": ws.max_row - 1, "deleted_excel_rows": sorted(drop), "kept": "last occurrence per date"}
for r in drop: ws.delete_rows(r)
dates = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
rep.update(rows_after=ws.max_row - 1, strictly_increasing=all(b > a_ for a_, b in zip(dates, dates[1:])),
           last_date=str(dates[-1].date()))
wb.save(a.xlsx)
if a.report: json.dump(rep, open(a.report, "w"), indent=2)
print(json.dumps(rep))
