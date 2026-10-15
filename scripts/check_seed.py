#!/usr/bin/env python3
"""Validate a completed raw_data.xlsx BEFORE deploying it as seed/raw_data.xlsx.  usage: check_seed.py seed/raw_data.xlsx
Errors (exit 1): wrong/duplicate header, dates not strictly increasing, duplicate dates, weekend rows, missing weekdays, empty cells,
text other than NA / numbers.  Warnings: absurd daily jumps, non-positive prices, NA in the last row, how far behind today the data is."""
import datetime as dt, json, math, os, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from ingest.common import WIB, load_config  # noqa: E402
from ingest.merge import is_equity  # noqa: E402


def check(path, cfg=None, today=None):
    cfg = cfg or load_config(); ds = cfg["dataset"]
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    err, warn, info = [], [], {}
    hdr = list(rows[0]); cols = hdr[1:]; body = rows[1:]
    if hdr[0] != ds["date_column"]:
        err.append(f"first header must be '{ds['date_column']}' (is {hdr[0]!r})")
    dup_h = sorted({c for c in cols if cols.count(c) > 1})
    if dup_h: err.append(f"duplicate column names: {dup_h}")
    missing_cfg = [k for k in cfg["instruments"] if k not in cols]
    if missing_cfg: err.append(f"columns required by config/pipeline.json are missing: {missing_cfg}")
    unknown = [c for c in cols if c not in cfg["instruments"] and not is_equity(c, ds["stock_pattern"])]
    if unknown: err.append(f"columns with no source rule (neither configured nor '<CODE> IJ EQUITY'): {unknown}")
    odd = [c for c in cols if c != str(c).strip() or "  " in str(c)]
    if odd: err.append(f"column names with stray spaces: {odd}")
    dates = []
    for i, r in enumerate(body, start=2):
        d = r[0].date() if isinstance(r[0], dt.datetime) else None
        if d is None: err.append(f"row {i}: first cell is not a date ({r[0]!r})"); continue
        dates.append((i, d))
    for (i1, a), (i2, b) in zip(dates, dates[1:]):
        if b == a: err.append(f"rows {i1}/{i2}: duplicate date {a}")
        elif b < a: err.append(f"rows {i1}/{i2}: dates go backwards ({a} -> {b})")
    wk = [(i, d) for i, d in dates if d.weekday() >= 5]
    if wk: err.append(f"weekend dates present (rows {[i for i, _ in wk][:5]}): the history uses weekdays only")
    if dates:
        have = {d for _, d in dates}; d0, d1 = dates[0][1], dates[-1][1]
        allwd = [d0 + dt.timedelta(days=k) for k in range((d1 - d0).days + 1) if (d0 + dt.timedelta(days=k)).weekday() < 5]
        gaps = [d for d in allwd if d not in have]
        if gaps: err.append(f"{len(gaps)} weekday(s) missing between {d0} and {d1}, e.g. {[str(g) for g in gaps[:8]]} (history has every weekday; holidays carry the last value)")
        info["first_date"], info["last_date"], info["rows"] = str(d0), str(d1), len(dates)
        t = today or dt.datetime.now(WIB).date()
        behind = [d for d in allwd + [d1 + dt.timedelta(days=k) for k in range(1, (t - d1).days + 1)] if d > d1 and d.weekday() < 5 and d <= t]
        info["weekdays_behind_today"] = len(behind)
        if behind: warn.append(f"data ends {d1}; {len(behind)} weekday(s) up to {t} will be filled by the pipeline")
    n_empty = n_text = n_num_text = 0; bad_text = []
    for i, r in enumerate(body, start=2):
        for j, v in enumerate(r[1:], start=1):
            if v is None: n_empty += 1; (len(bad_text) < 6 and bad_text.append(f"{hdr[j]}@row{i}: EMPTY (must be the text NA)"))
            elif isinstance(v, str):
                if v == "NA": n_text += 1
                else:
                    try: float(v); n_num_text += 1
                    except ValueError: bad_text.append(f"{hdr[j]}@row{i}: text {v!r} (only NA or numbers allowed)")
            elif not isinstance(v, (int, float)) or (isinstance(v, float) and math.isnan(v)): bad_text.append(f"{hdr[j]}@row{i}: unusable value {v!r}")
    if n_empty: err.append(f"{n_empty} empty cell(s): missing values must be the literal text NA (history has 0 empty cells)")
    err += [x for x in bad_text if "EMPTY" not in x][:8]
    info["NA_cells"], info["numbers_stored_as_text"] = n_text, n_num_text
    if n_num_text: warn.append(f"{n_num_text} number(s) stored as text (R as.numeric handles them)")
    big, nonpos = [], []
    for j, c in enumerate(cols, start=1):
        prev = None
        for i, r in enumerate(body, start=2):
            v = r[j]
            try: x = float(v)
            except (TypeError, ValueError): prev = None if v == "NA" else prev; continue
            if x <= 0 and c != "GTIDR1Y Govt": nonpos.append(f"{c}@row{i}={x}")
            elif prev and prev > 0 and x > 0 and abs(math.log(x / prev)) > ds["max_abs_daily_logret"]: big.append(f"{c}@row{i}: {prev}->{x}")
            prev = x
    if nonpos: warn.append(f"non-positive values: {nonpos[:6]}")
    if big: warn.append(f"daily jumps beyond {ds['max_abs_daily_logret']} (typo? corporate action?): {big[:6]}")
    phei = Path(path).with_name("phei_history.csv")
    if dates:
        from ingest.phei_store import load as phei_load
        have = phei_load(phei) if phei.exists() else {}
        if dates[-1][1] not in have:
            warn.append(f"IBPRCPI: {phei.name} has no PHEI value for the last data date {dates[-1][1]}. PHEI shows only the last 2 days, so add a line "
                        f"'{dates[-1][1]},<value>,last' (value from phei.co.id) or the first run writes NA for IBPRCPI until the capture supplies it.")
    if body:
        na_last = [hdr[j] for j, v in enumerate(body[-1][1:], start=1) if v == "NA" and hdr[j] in cfg["instruments"]]
        if na_last: warn.append(f"required instruments are NA on the last row: {na_last}")
    return err, warn, info


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "seed/raw_data.xlsx"
    err, warn, info = check(path)
    md = [f"## Seed check: {'FAILED' if err else 'OK'}  ({path})", "", f"`{json.dumps(info)}`", ""]
    md += ["### Errors (must fix)"] + [f"- {e}" for e in err] if err else ["No errors."]
    if warn: md += ["", "### Warnings"] + [f"- {w}" for w in warn]
    out = "\n".join(md); print(out)
    if os.environ.get("GITHUB_STEP_SUMMARY"): open(os.environ["GITHUB_STEP_SUMMARY"], "a").write(out + "\n")
    sys.exit(1 if err else 0)
