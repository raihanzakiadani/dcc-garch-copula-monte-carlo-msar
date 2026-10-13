#!/usr/bin/env python3
"""Readable Markdown for a data-only test: what was fetched, what was appended, what failed. usage: ingest_summary.py work"""
import json, os, sys
from pathlib import Path
from openpyxl import load_workbook

w = Path(sys.argv[1]); r = json.loads((w / "ingest_result.json").read_text())
md = [f"## Data test: ingest status **{r['status']}**", "",
      f"- last date before: `{r.get('last_date_before')}` | after: `{r.get('last_date_after')}` | row limit (newest finalised weekday): `{r.get('row_limit')}`",
      f"- new rows appended: **{len(r.get('new_dates') or [])}** {r.get('new_dates') or ''}", f"- reason: {r.get('reason') or '-'}",
      f"- blocked dates: {r.get('blocked_dates') or '-'}", f"- revisions seen vs stored: {len(r.get('revisions') or [])}", ""]
md += ["| instrument | status | new values | carried | detail / definition check |", "|---|---|---|---|---|"]
for k, v in r["instruments"].items():
    if k.endswith("IJ EQUITY") and not str(v["status"]).startswith("FAILED"):
        continue                                        # list stocks only when they failed
    d = v.get("detail") or (", ".join(f"{a}={b}" for a, b in (v.get("diag") or {}).items()) if v.get("diag") else "")
    md.append(f"| {k} | {v['status']} | {v.get('n_new_values', '')} | {v.get('n_carried', '')} | {d} |")
ok_eq = sum(1 for k, v in r["instruments"].items() if k.endswith("IJ EQUITY") and not str(v["status"]).startswith("FAILED"))
n_eq = sum(1 for k in r["instruments"] if k.endswith("IJ EQUITY"))
md.append(f"| {n_eq} IDX stocks | {ok_eq} not failed | | | failed ones are listed above |")
if r.get("warnings"):
    md += ["", "### Warnings"] + [f"- {x}" for x in r["warnings"]]
x = w / "raw_data.xlsx"
if x.exists() and r.get("new_dates"):
    ws = load_workbook(x, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
    show = ["Dates", "JCI Index", "USDIDR BGN Curncy", "CO1 Comdty", "XAU Curncy", "XW1 Comdty", "KO1 Comdty", "LMNIDS03 LME Comdty", "IBPRCPI Index", "GTIDR1Y Govt", "BBCA IJ EQUITY"]
    idx = [hdr.index(c) for c in show if c in hdr]
    md += ["", "### Last rows of raw_data.xlsx (previous stored row first, then the new rows)", "",
           "| " + " | ".join(show) + " |", "|" + "---|" * len(show)]
    n_new = len(r["new_dates"])
    for row in rows[-(n_new + 1):]:
        md.append("| " + " | ".join(str(row[i].date()) if hasattr(row[i], "date") else str(row[i]) for i in idx) + " |")
s = "\n".join(md); print(s)
if os.environ.get("GITHUB_STEP_SUMMARY"):
    open(os.environ["GITHUB_STEP_SUMMARY"], "a").write(s + "\n")
